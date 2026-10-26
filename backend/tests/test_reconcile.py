"""Reconciliation: match two tables, and explain every row that does not agree."""
import io

import pandas as pd
import pytest

from app import reconcile
from app.config import settings

CFG = {"left": "courier", "right": "bank", "left_key": "waybill", "right_key": "ref",
       "left_amount": "cod", "right_amount": "paid", "tolerance": 0.01}


def frames(left_rows, right_rows):
    return (pd.DataFrame(left_rows, columns=["waybill", "cod"]), pd.DataFrame(right_rows, columns=["ref", "paid"]))


def kinds(out):
    return sorted(out["kind"])


# ---------- the matching rules

def test_everything_agrees():
    out, m, d = reconcile.compare(*frames([["A1", 100], ["A2", 50.5]], [["A1", 100], ["A2", 50.5]]), CFG)
    assert len(out) == 0 and m["matched_rows"] == 2 and m["exceptions"] == 0
    assert d["matched_amount"] == 150.5 and reconcile.summary_line(m, "courier", "bank").endswith("No exceptions.")


def test_each_kind_of_problem_is_found_with_a_reason():
    left = [["A1", 100], ["A2", 200], ["A3", 300], ["A4", None]]
    right = [["A1", 100], ["A2", 190], ["A4", 40], ["A9", 77]]
    out, m, _ = reconcile.compare(*frames(left, right), CFG)
    by = {r.key: r for r in out.itertuples()}
    assert by["A2"].kind == "amount_mismatch" and by["A2"].difference == 10.0 and "difference +10.00" in by["A2"].reason
    assert by["A3"].kind == "missing_in_right" and by["A3"].left_amount == 300 and "not in 'bank'" in by["A3"].reason
    assert by["A9"].kind == "missing_in_left" and by["A9"].right_amount == 77 and "not in 'courier'" in by["A9"].reason
    assert by["A4"].kind == "unreadable_amount" and "empty or unreadable in 'courier'" in by["A4"].reason
    assert m["matched_rows"] == 1 and m["exceptions"] == 4
    assert m["amount_difference"] == 10.0 and m["unmatched_amount"] == 377.0  # 300 only in courier + 77 only in bank


def test_a_small_difference_inside_the_tolerance_still_matches():
    cfg = {**CFG, "tolerance": 0.5}
    out, m, _ = reconcile.compare(*frames([["A1", 100.00], ["A2", 100.00]], [["A1", 100.40], ["A2", 100.60]]), cfg)
    assert m["matched_rows"] == 1 and list(out["key"]) == ["A2"]
    strict, m2, _ = reconcile.compare(*frames([["A1", 100.00]], [["A1", 100.40]]), {**CFG, "tolerance": 0})
    assert m2["amount_mismatches"] == 1


def test_keys_ignore_capitals_spaces_and_excel_float_endings():
    left = [[" ab-1 ", 10], ["x  y", 20], [1001, 30]]
    right = [["AB-1", 10], ["X Y", 20], [1001.0, 30]]
    out, m, _ = reconcile.compare(*frames(left, right), CFG)
    assert m["matched_rows"] == 3 and len(out) == 0


def test_amounts_in_text_form_are_understood():
    out, m, _ = reconcile.compare(*frames([["A1", "SAR 1,200.50"], ["A2", "(75.00)"]], [["A1", "1200.5"], ["A2", -75]]), CFG)
    assert m["matched_rows"] == 2 and len(out) == 0


def test_duplicate_keys_are_flagged_and_the_first_row_is_used_for_matching():
    out, m, _ = reconcile.compare(*frames([["A1", 100], ["A1", 100], ["A2", 5]], [["A1", 100], ["A2", 5]]), CFG)
    assert m["duplicate_keys"] == 1 and m["matched_rows"] == 2 and m["missing_in_left"] == 0
    row = out.iloc[0]
    assert row.kind == "duplicate" and row.left_row == 3 and "Row 2 was used for matching, row 3 is the extra" in row.reason


def test_rows_with_no_key_are_listed_not_dropped():
    out, m, _ = reconcile.compare(*frames([[None, 10], ["A1", 5]], [["A1", 5], ["", 3]]), CFG)
    assert kinds(out) == ["missing_key", "missing_key"] and m["matched_rows"] == 1
    assert "Row 2 in 'courier' has no key" in out.iloc[0].reason


def test_without_amount_columns_only_the_keys_are_compared():
    cfg = {k: v for k, v in CFG.items() if "amount" not in k}
    out, m, d = reconcile.compare(*frames([["A1", 1], ["A2", 2]], [["A1", 999]]), cfg)
    assert m["matched_rows"] == 1 and kinds(out) == ["missing_in_right"] and d["amounts_compared"] is False


def test_row_numbers_match_the_spreadsheet_row():
    out, _, _ = reconcile.compare(*frames([["A1", 1], ["A2", 2]], [["A1", 1]]), CFG)
    assert out.iloc[0].left_row == 3  # header is row 1, A1 is row 2, A2 is row 3


def test_worst_problems_come_first():
    left = [["B", 1], ["A", 2], ["C", 3]]
    right = [["B", 9], ["Z", 4]]
    out, _, _ = reconcile.compare(*frames(left, right), CFG)
    assert list(out["kind"]) == ["amount_mismatch", "missing_in_right", "missing_in_right", "missing_in_left"]


# ---------- through the API

def add(client, name, text):
    r = client.post("/api/datasources/uploads", files={"file": (name, io.BytesIO(text.encode()), "text/csv")})
    assert r.status_code == 201, r.text
    return r.json()["name"]


def recon_steps(a, b, **over):
    cfg = {"left": "courier", "right": "bank", "left_key": "waybill", "right_key": "ref", "left_amount": "cod", "right_amount": "paid", **over}
    return [
        {"id": "s1", "type": "trigger", "name": "Start", "config": {}},
        {"id": "s2", "type": "load_csv", "name": "Courier collections", "config": {"file": a, "source": "uploads", "table": "courier"}},
        {"id": "s3", "type": "load_csv", "name": "Bank deposits", "config": {"file": b, "source": "uploads", "table": "bank"}},
        {"id": "s4", "type": "compare", "name": "Match", "config": cfg},
        {"id": "s5", "type": "report", "name": "Report", "config": {"table": "courier"}},
        {"id": "s6", "type": "end", "name": "Done", "config": {}},
    ]


@pytest.fixture(scope="module")
def run_ctx(client):
    a = add(client, "courier.csv", "waybill,cod\nW1,100\nW2,200\nW3,300\nW4,50\nW4,50\n")
    b = add(client, "bank.csv", "ref,paid\nW1,100\nW2,190\nW9,70\nW4,50\n")
    wf = client.post("/api/workflows", json={"name": "Recon API test", "steps": recon_steps(a, b)})
    assert wf.status_code == 201, wf.text
    ex = client.post(f"/api/workflows/{wf.json()['id']}/run").json()
    return ex


def test_the_run_succeeds_and_says_what_it_found(client, run_ctx):
    assert run_ctx["status"] == "success"
    msg = next(s for s in run_ctx["steps"] if s["step_id"] == "s4")["message"]
    assert msg.startswith("Matched 2 of 4 rows from 'courier'") and "1 amount differences" in msg and "1 duplicates" in msg


def test_the_run_keeps_its_numbers_after_it_ends(client, run_ctx):
    got = client.get(f"/api/executions/{run_ctx['id']}").json()["result"]
    assert got["metrics"]["matched_rows"] == 2 and got["metrics"]["exceptions"] == 4 and got["metrics"]["amount_difference"] == 10.0
    cmp = got["details"]["comparisons"][0]
    assert cmp["left"] == "courier" and cmp["counts"]["missing_in_right"] == 1 and cmp["net_difference"] == 10.0


def test_exceptions_can_be_listed_filtered_and_paged(client, run_ctx):
    base = f"/api/executions/{run_ctx['id']}/exceptions"
    d = client.get(base).json()
    assert d["total"] == 4 and d["by_kind"] == {"amount_mismatch": 1, "missing_in_right": 1, "missing_in_left": 1, "duplicate": 1}
    assert [i["kind"] for i in d["items"]][0] == "amount_mismatch" and d["labels"]["duplicate"] == "Duplicate key"
    one = client.get(base, params={"kind": "missing_in_left"}).json()
    assert one["total"] == 1 and one["items"][0]["key"] == "W9" and one["items"][0]["decision"] is None
    assert len(client.get(base, params={"limit": 1, "offset": 3}).json()["items"]) == 1
    assert client.get(base, params={"kind": "nonsense"}).status_code == 422
    assert client.get(base, params={"limit": 0}).status_code == 422
    assert client.get("/api/executions/99999/exceptions").status_code == 404


def test_the_exceptions_csv_download_is_safe_for_excel(client, run_ctx):
    r = client.get(f"/api/executions/{run_ctx['id']}/download/exceptions")
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    df = pd.read_csv(io.StringIO(r.text))
    assert len(df) == 4 and {"type", "key", "reason", "decision"} <= set(df.columns)
    assert client.get("/api/executions/1/download/exceptions").status_code == 404  # a run with no exceptions


def test_the_excel_report_has_an_exceptions_sheet(client, run_ctx):
    import openpyxl
    wb = openpyxl.load_workbook(settings.reports_dir / f"execution_{run_ctx['id']}_report.xlsx")
    assert "Exceptions" in wb.sheetnames
    rows = list(wb["Exceptions"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4 and rows[0][0] == "Amount differs"


def test_a_condition_can_branch_on_the_reconciliation_result(client):
    a = add(client, "c2.csv", "waybill,cod\nW1,100\nW2,5\n")
    b = add(client, "b2.csv", "ref,paid\nW1,100\n")
    steps = recon_steps(a, b)[:4] + [
        {"id": "s5", "type": "condition", "name": "Any exceptions?", "config": {"field": "exceptions", "op": ">", "value": 0, "on_true": "s6", "on_false": "s7"}},
        {"id": "s6", "type": "notify", "name": "Alert finance", "next": "s7", "config": {"to": "finance@example.com", "message": "Differences found"}},
        {"id": "s7", "type": "end", "name": "Done", "config": {}}]
    wf = client.post("/api/workflows", json={"name": "Recon branch test", "steps": steps})
    assert wf.status_code == 201, wf.text
    ex = client.post(f"/api/workflows/{wf.json()['id']}/run").json()
    assert ex["status"] == "success" and "Alert finance" in [s["name"] for s in ex["steps"]]


def test_a_wrong_column_name_fails_the_run_and_lists_the_real_columns(client):
    a = add(client, "c3.csv", "waybill,cod\nW1,100\n")
    b = add(client, "b3.csv", "ref,paid\nW1,100\n")
    wf = client.post("/api/workflows", json={"name": "Recon bad column", "steps": recon_steps(a, b, right_key="reference")})
    ex = client.post(f"/api/workflows/{wf.json()['id']}/run").json()
    assert ex["status"] == "failed" and "'reference' is not in table 'bank'. Columns there: ref, paid" in ex["error_message"]


def test_the_editor_rejects_a_badly_set_up_compare_step(client):
    ok = recon_steps("a.csv", "b.csv")
    def post(**over):
        steps = [dict(s) for s in ok]
        steps[3] = {**steps[3], "config": {**steps[3]["config"], **over}}
        return client.post("/api/workflows", json={"name": f"Bad compare {len(over)}{list(over.values())}", "steps": steps})
    assert post(right="courier").status_code == 422        # the same table twice
    assert post(right="ghost").status_code == 422          # a table nobody loads
    assert post(left_key="").status_code == 422
    assert post(right_amount="").status_code == 422        # one amount column only
    assert post(tolerance=-1).status_code == 422
    assert post(tolerance=True).status_code == 422
    assert post(tolerance="a lot").status_code == 422
    assert post(tolerance=0).status_code == 201
