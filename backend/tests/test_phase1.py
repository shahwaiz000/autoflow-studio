"""Phase 1: Excel files, several named tables, files chosen at run time, messy-value cleaning, real reports."""
import io
import time
from datetime import timedelta

import numpy as np
import openpyxl
import pandas as pd
import pytest

from app import tables
from app.config import settings
from app.database import SessionLocal
from app.models import utcnow
from app import schedules


# ---------- helpers

def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def add_file(client, name, data):
    r = client.post("/api/datasources/uploads", files={"file": (name, io.BytesIO(data), "application/octet-stream")})
    assert r.status_code == 201, r.text
    return r.json()


def make_wf(client, name, steps):
    r = client.post("/api/workflows", json={"name": name, "steps": steps})
    assert r.status_code == 201, r.text
    return r.json()


def T(id_, name="Start"):
    return {"id": id_, "type": "trigger", "name": name, "config": {}}


def run(client, wid, inputs=None, **params):
    return client.post(f"/api/workflows/{wid}/run", params=params, json={"inputs": inputs} if inputs is not None else None)


def wait(client, eid):
    end = time.time() + 10
    while time.time() < end:
        ex = client.get(f"/api/executions/{eid}").json()
        if ex["status"] not in ("queued", "running"):
            return ex
        time.sleep(0.05)
    raise AssertionError("run did not finish")


# ---------- cleaning messy values

@pytest.mark.parametrize("text,expected", [
    ("1,234.50 SAR", 1234.5), ("SAR 1.234,50", 1234.5), ("(500.00)", -500.0), ("12,50", 12.5),
    ("1,234", 1234.0), ("-45", -45.0), (" 99 ", 99.0), (7, 7.0), ("1.234.567", 1234567.0)])
def test_amounts_are_read_the_way_a_person_reads_them(text, expected):
    assert tables.to_amount(text) == expected


@pytest.mark.parametrize("text", ["abc", "", None, "N/A"])
def test_unreadable_amounts_become_empty_not_zero(text):
    assert np.isnan(tables.to_amount(text))


def test_normalize_amounts_counts_what_it_fixed_and_what_it_could_not_read():
    out, fixed, bad = tables.normalize_amounts(pd.Series(["1,000.00 SAR", "250", "oops", None, 12.5]))
    assert list(out.dropna()) == [1000.0, 250.0, 12.5] and fixed == 2 and bad == 1


def test_dates_become_one_format_and_day_first_is_the_default():
    out, changed, bad = tables.parse_dates(pd.Series(["05/10/2026", "2026-10-06", "not a date", None]))
    assert out.iloc[0] == "2026-10-05" and out.iloc[1] == "2026-10-06"  # 5 October, not 10 May
    assert changed == 1 and bad == 1
    month_first, *_ = tables.parse_dates(pd.Series(["05/10/2026"]), dayfirst=False)
    assert month_first.iloc[0] == "2026-05-10"
    for day_first in (True, False):  # a date already in ISO form must never be flipped
        kept, *_ = tables.parse_dates(pd.Series(["2026-10-06", "2026-12-01"]), dayfirst=day_first)
        assert list(kept) == ["2026-10-06", "2026-12-01"]
    stamp, *_ = tables.parse_dates(pd.Series([pd.Timestamp("2026-10-06")]))  # real Excel dates
    assert stamp.iloc[0] == "2026-10-06"


def test_text_is_trimmed_collapsed_and_upper_cased():
    out, changed = tables.standardize_text(pd.Series(["  inv-1  ", "inv  2", "INV-3", None]))
    assert list(out.dropna()) == ["INV-1", "INV 2", "INV-3"] and changed == 2


# ---------- reading files

def test_excel_sheets_can_be_listed_and_chosen(tmp_path):
    p = tmp_path / "book.xlsx"
    p.write_bytes(xlsx_bytes({"Jan": [["id", "amt"], [1, 10]], "Feb": [["id", "amt"], [2, 20], [3, 30]]}))
    assert tables.sheet_names(p) == ["Jan", "Feb"]
    assert len(tables.read_table(p)) == 1 and len(tables.read_table(p, sheet="Feb")) == 2
    with pytest.raises(tables.TableError, match="Sheets in this file: Jan, Feb"):
        tables.read_table(p, sheet="Mar")


def test_a_fake_xlsx_is_refused_with_a_plain_message(tmp_path):
    p = tmp_path / "fake.xlsx"
    p.write_bytes(b"this is just text")
    with pytest.raises(tables.TableError, match="not a real .xlsx"):
        tables.read_table(p)


# ---------- Data Sources with Excel

def test_excel_upload_is_listed_with_its_sheets_and_can_be_previewed(client):
    f = add_file(client, "q3 payments.xlsx", xlsx_bytes({"Payments": [["ref", "amount"], ["A1", 100], ["A2", 250]], "Notes": [["x"], ["y"]]}))
    row = next(u for u in client.get("/api/datasources").json()["uploads"] if u["name"] == f["name"])
    assert row["original_name"] == "q3 payments.xlsx" and row["rows"] == 2 and row["sheets"] == ["Payments", "Notes"]
    p = client.get(f"/api/datasources/uploads/{f['name']}/preview").json()
    assert p["columns"] == ["ref", "amount"] and p["rows"][0] == ["A1", "100"]
    other = client.get(f"/api/datasources/uploads/{f['name']}/preview", params={"sheet": "Notes"}).json()
    assert other["columns"] == ["x"]


@pytest.mark.parametrize("name,data,message", [
    ("notes.txt", b"hello", "Only .csv and .xlsx"),
    ("empty.csv", b"   ", "empty"),
    ("header_only.csv", b"a,b\n", "no data rows"),
    ("fake.xlsx", b"not a zip", "not a real .xlsx"),
])
def test_bad_uploads_are_refused_and_leave_no_file_behind(client, name, data, message):
    before = set(p.name for p in settings.uploads_dir.glob("*")) if settings.uploads_dir.exists() else set()
    r = client.post("/api/datasources/uploads", files={"file": (name, io.BytesIO(data), "x")})
    assert r.status_code == 400 and message in r.json()["error"]["message"]
    after = set(p.name for p in settings.uploads_dir.glob("*")) if settings.uploads_dir.exists() else set()
    assert after == before


def test_the_cleaner_upload_page_accepts_excel_too(client):
    data = xlsx_bytes({"S": [["name", "amount"], ["a", 1], ["a", 1], ["b", 2]]})
    r = client.post("/api/uploads", files={"file": ("sheet.xlsx", io.BytesIO(data), "x")})
    assert r.status_code == 201 and r.json()["status"] == "success"


# ---------- several tables in one workflow

def two_tables(extra_steps=()):
    return [
        T("s1"),
        {"id": "s2", "type": "load_csv", "name": "Load invoices", "config": {"file": "invoices.csv", "table": "invoices"}},
        {"id": "s3", "type": "load_csv", "name": "Load customers", "config": {"file": "customers.csv", "table": "customers"}},
        *extra_steps,
    ]


def test_two_files_are_held_at_once_and_each_step_uses_the_table_it_names(client):
    steps = two_tables([
        {"id": "s4", "type": "transform", "name": "Clean customers", "config": {"operations": ["drop_duplicates"], "table": "customers"}},
        {"id": "s5", "type": "validate", "name": "Check invoices", "config": {"table": "invoices"}},
    ])
    ex = run(client, make_wf(client, "Two tables", steps)["id"]).json()
    assert ex["status"] == "success"
    by = {s["step_id"]: s["message"] for s in ex["steps"]}
    assert "into table 'invoices'" in by["s2"] and "into table 'customers'" in by["s3"]
    assert by["s4"].startswith("Applied drop_duplicates: 43 ->")  # customers.csv has 43 rows
    assert "of 40 rows" in by["s5"] or "rows have missing values" in by["s5"]


def test_a_step_that_names_a_table_that_is_not_loaded_fails_clearly(client):
    r = client.post("/api/workflows", json={"name": "Ghost table", "steps": [T("s1"), {
        "id": "s2", "type": "validate", "name": "Check", "config": {"table": "ghost"}}]})
    assert r.status_code == 422 and "no Load step creates it" in r.json()["error"]["message"]


def test_duplicate_table_names_and_bad_table_names_are_refused(client):
    load = lambda i, t: {"id": i, "type": "load_csv", "name": f"Load {i}", "config": {"file": "invoices.csv", "table": t}}
    r = client.post("/api/workflows", json={"name": "Dup tables", "steps": [T("s1"), load("s2", "bank"), load("s3", "bank")]})
    assert r.status_code == 422 and "own name" in r.json()["error"]["message"]
    r = client.post("/api/workflows", json={"name": "Bad table", "steps": [T("s1"), load("s2", "Bank Data!")]})
    assert r.status_code == 422


# ---------- files chosen at run time

def input_wf(client, name="Needs a file", sheet=None, table="main"):
    cfg = {"input": "bank", "table": table, **({"sheet": sheet} if sheet else {})}
    return make_wf(client, name, [T("s1"), {"id": "s2", "type": "load_csv", "name": "Bank statement", "config": cfg},
                                  {"id": "s3", "type": "validate", "name": "Check", "config": {}}])


def test_a_workflow_lists_the_files_it_asks_for(client):
    wf = input_wf(client)
    assert wf["inputs"] == [{"name": "bank", "label": "Bank statement"}]
    listed = next(w for w in client.get("/api/workflows").json() if w["id"] == wf["id"])
    assert listed["inputs"] == wf["inputs"]
    assert next(w for w in client.get("/api/workflows").json() if w["name"] == "Invoice Processing")["inputs"] == []


def test_running_without_choosing_a_file_is_refused_with_the_name_of_the_file(client):
    wid = input_wf(client, "Needs a file 2")["id"]
    r = run(client, wid)
    assert r.status_code == 422 and "Choose a file for 'Bank statement'" in r.json()["error"]["message"]


def test_unknown_or_missing_or_unsafe_choices_are_refused(client):
    wid = input_wf(client, "Needs a file 3")["id"]
    assert run(client, wid, {"nope": {"source": "samples", "file": "invoices.csv"}}).status_code == 422
    assert run(client, wid, {"bank": {"source": "samples", "file": "missing.csv"}}).status_code == 422
    assert run(client, wid, {"bank": {"source": "samples", "file": "../x.csv"}}).status_code == 422
    assert run(client, wid, {"bank": {"source": "elsewhere", "file": "invoices.csv"}}).status_code == 422


def test_a_workflow_with_no_file_slots_rejects_files_it_did_not_ask_for(client):
    wid = next(w["id"] for w in client.get("/api/workflows").json() if w["name"] == "Invoice Processing")
    assert run(client, wid, {"bank": {"source": "samples", "file": "invoices.csv"}}).status_code == 422


def test_run_on_the_chosen_file_and_remember_which_file_it_was(client):
    wid = input_wf(client, "Needs a file 4")["id"]
    ex = run(client, wid, {"bank": {"source": "samples", "file": "sales_orders.csv"}}).json()
    assert ex["status"] == "success" and ex["inputs"] == [{"name": "bank", "source": "samples", "file": "sales_orders.csv"}]
    assert "sales_orders.csv" in next(s for s in ex["steps"] if s["step_id"] == "s2")["message"]
    again = client.get(f"/api/executions/{ex['id']}").json()
    assert again["inputs"] == ex["inputs"]  # the record survives


def test_the_same_workflow_runs_on_different_files(client):
    wid = input_wf(client, "Needs a file 5")["id"]
    a = run(client, wid, {"bank": {"source": "samples", "file": "invoices.csv"}}).json()
    b = run(client, wid, {"bank": {"source": "samples", "file": "employees.csv"}}).json()
    rows = lambda e: next(s for s in e["steps"] if s["step_id"] == "s2")["message"]
    assert rows(a) != rows(b) and a["inputs"][0]["file"] != b["inputs"][0]["file"]


def test_background_runs_carry_the_chosen_files_to_the_worker(client):
    wid = input_wf(client, "Needs a file 6")["id"]
    r = run(client, wid, {"bank": {"source": "samples", "file": "customers.csv"}}, background="true")
    assert r.status_code == 202 and r.json()["inputs"][0]["file"] == "customers.csv"
    done = wait(client, r.json()["id"])
    assert done["status"] == "success" and "customers.csv" in next(s for s in done["steps"] if s["step_id"] == "s2")["message"]


def test_an_uploaded_excel_sheet_can_be_chosen_at_run_time(client):
    f = add_file(client, "bank.xlsx", xlsx_bytes({"Cover": [["note"], ["x"]], "Txns": [["ref", "amount"], ["A", 1], ["B", 2], ["C", 3]]}))
    wid = input_wf(client, "Needs a file 7", sheet="Txns")["id"]
    ex = run(client, wid, {"bank": {"source": "uploads", "file": f["name"]}}).json()
    assert ex["status"] == "success" and "Loaded 3 rows and 2 columns" in next(s for s in ex["steps"] if s["step_id"] == "s2")["message"]


def test_a_missing_sheet_fails_the_run_with_the_sheet_names(client):
    f = add_file(client, "bank2.xlsx", xlsx_bytes({"Txns": [["ref"], ["A"]]}))
    wid = input_wf(client, "Needs a file 8", sheet="Nope")["id"]
    ex = run(client, wid, {"bank": {"source": "uploads", "file": f["name"]}}).json()
    assert ex["status"] == "failed" and "Sheets in this file: Txns" in ex["error_message"]


def test_scheduled_runs_skip_workflows_that_need_files_and_say_why(client):
    wid = input_wf(client, "Needs a file 9")["id"]
    sid = client.post("/api/schedules", json={"workflow_id": wid, "kind": "interval", "interval_minutes": 5}).json()["id"]
    with SessionLocal() as db:
        ran = [x for x in schedules.run_due(db, utcnow() + timedelta(minutes=6)) if x["schedule_id"] == sid]
    assert ran[0]["status"] == "Skipped: this workflow needs files chosen when it runs" and ran[0]["execution_id"] is None


# ---------- tidy-up steps

def test_messy_amounts_and_dates_are_cleaned_and_counted(client):
    csv = ("ref,amount,date,vendor\n"
           "A1,\"SAR 1,000.00\",05/10/2026,  acme  ltd\n"
           "A2,250,2026-10-06,Acme Ltd\n"
           "A3,oops,31/12/2026,Beta\n").encode()
    f = add_file(client, "messy.csv", csv)
    steps = [T("s1"),
             {"id": "s2", "type": "load_csv", "name": "Load", "config": {"file": f["name"], "source": "uploads"}},
             {"id": "s3", "type": "transform", "name": "Tidy", "config": {
                 "operations": ["normalize_amounts", "parse_dates", "standardize_text"],
                 "amount_columns": ["amount"], "date_columns": ["date"], "text_columns": ["vendor"]}},
             {"id": "s4", "type": "report", "name": "Report", "config": {}}]
    ex = run(client, make_wf(client, "Messy data", steps)["id"]).json()
    assert ex["status"] == "success"
    msg = next(s for s in ex["steps"] if s["step_id"] == "s3")["message"]
    assert "could not be read and were left empty" in msg and "reformatted" in msg
    out = pd.read_csv(settings.reports_dir / f"execution_{ex['id']}_processed.csv")
    assert list(out["amount"].dropna()) == [1000.0, 250.0] and out["amount"].isna().sum() == 1
    assert list(out["date"]) == ["2026-10-05", "2026-10-06", "2026-12-31"]
    assert list(out["vendor"]) == ["ACME LTD", "ACME LTD", "BETA"]


def test_a_cleaning_option_must_say_which_columns(client):
    steps = [T("s1"), {"id": "s2", "type": "load_csv", "name": "Load", "config": {"file": "invoices.csv"}},
             {"id": "s3", "type": "transform", "name": "Tidy", "config": {"operations": ["normalize_amounts"]}}]
    r = client.post("/api/workflows", json={"name": "No columns", "steps": steps})
    assert r.status_code == 422 and "name the columns" in r.json()["error"]["message"]


def test_a_cleaning_step_on_a_missing_column_fails_clearly(client):
    steps = [T("s1"), {"id": "s2", "type": "load_csv", "name": "Load", "config": {"file": "invoices.csv"}},
             {"id": "s3", "type": "transform", "name": "Tidy", "config": {"operations": ["parse_dates"], "date_columns": ["nope"]}}]
    ex = run(client, make_wf(client, "Missing col", steps)["id"]).json()
    assert ex["status"] == "failed" and "Missing required columns: nope" in ex["error_message"]


# ---------- the report

def test_the_report_is_a_real_workbook_with_summary_columns_and_data(client):
    wid = next(w["id"] for w in client.get("/api/workflows").json() if w["name"] == "Sales Report Generation")
    ex = run(client, wid).json()
    assert ex["status"] == "success"
    wb = openpyxl.load_workbook(settings.reports_dir / f"execution_{ex['id']}_report.xlsx")
    assert wb.sheetnames == ["Summary", "Columns", "Data"]
    cols = {r[0]: r for r in wb["Columns"].iter_rows(min_row=2, values_only=True)}
    assert cols["revenue"][1] == "number" and cols["revenue"][6] > 0 and cols["region"][1] == "text"
    facts = dict(wb["Summary"].iter_rows(min_row=2, values_only=True))
    assert facts["Rows in the report"] > 0
    summary = pd.read_csv(settings.reports_dir / f"execution_{ex['id']}_summary.csv")
    assert {"column", "type", "missing", "total"} <= set(summary.columns)  # no longer just describe()


def test_report_text_that_looks_like_a_formula_is_neutralised_in_excel(client):
    f = add_file(client, "evil.csv", b"name,amount\n=HYPERLINK(\"http://evil.example\"),5\n+cmd,6\nplain,7\n")
    steps = [T("s1"), {"id": "s2", "type": "load_csv", "name": "Load", "config": {"file": f["name"], "source": "uploads"}},
             {"id": "s3", "type": "report", "name": "Report", "config": {}}]
    ex = run(client, make_wf(client, "Formula safety", steps)["id"]).json()
    wb = openpyxl.load_workbook(settings.reports_dir / f"execution_{ex['id']}_report.xlsx")
    cells = [r[0] for r in wb["Data"].iter_rows(min_row=2, values_only=True)]
    assert cells[0].startswith("'=") and cells[1].startswith("'+") and cells[2] == "plain"
    assert all(not c.startswith("=") for c in cells)


def test_the_workbook_can_be_downloaded(client):
    wid = next(w["id"] for w in client.get("/api/workflows").json() if w["name"] == "Sales Report Generation")
    ex = run(client, wid).json()
    r = client.get(f"/api/executions/{ex['id']}/download/report")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "spreadsheetml" in r.headers["content-type"]
