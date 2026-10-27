"""POS Collection Reconciliation: orders closed in the system against the drivers' POS collections (all fictional)."""
import pandas as pd
import pytest
from openpyxl import load_workbook

from app import reconcile
from app.config import settings
from app.definitions import clean_definition


def run_pos(client):
    wf = client.post("/api/templates/pos-reconciliation/use").json()
    files = {"system": {"source": "samples", "file": "pos_system_orders.csv"}, "pos": {"source": "samples", "file": "pos_collections.csv"}}
    return client.post(f"/api/workflows/{wf['id']}/run", json={"inputs": files}).json()


def test_the_planted_answer_is_found_exactly(client):
    ex = run_pos(client)
    m = ex["result"]["metrics"]
    assert ex["status"] == "success"
    assert (m["matched_rows"], m["exceptions"]) == (53, 10)
    assert (m["amount_mismatches"], m["missing_in_right"], m["missing_in_left"], m["duplicate_keys"]) == (5, 2, 2, 1)
    items = client.get(f"/api/executions/{ex['id']}/exceptions").json()["items"]
    diffs = {i["key"]: i["difference"] for i in items if i["kind"] == "amount_mismatch"}
    # difference = system amount minus POS amount: positive means the system recorded more than the POS collected
    assert diffs == {"ORD-50005": 18.0, "ORD-50016": -10.0, "ORD-50028": 90.0, "ORD-50039": 0.5, "ORD-50053": -5.0}
    assert {i["key"] for i in items if i["kind"] == "missing_in_right"} == {"ORD-50010", "ORD-50034"}
    assert {i["key"] for i in items if i["kind"] == "missing_in_left"} == {"ORD-59901", "ORD-59902"}
    assert [i["key"] for i in items if i["kind"] == "duplicate"] == ["ORD-50021"]


def test_untidy_order_ids_and_text_amounts_still_match(client):
    ex = run_pos(client)
    keys = {i["key"] for i in client.get(f"/api/executions/{ex['id']}/exceptions").json()["items"]}
    # ord-50004, " ORD-50006 " and "SAR 1,234.50" style amounts are in the file, and none of them is an exception
    pos = pd.read_csv(settings.samples_dir / "pos_collections.csv")
    assert (pos["order_id"].str.strip().str.upper() != pos["order_id"]).any() and pos["amount"].astype(str).str.startswith("SAR").any()
    assert "ORD-50004" not in keys and "ORD-50006" not in keys


def test_each_driver_gets_a_summary_that_adds_up_to_the_total(client):
    ex = run_pos(client)
    g = ex["result"]["details"]["comparisons"][0]["group_by"]
    assert g["column"] == "driver" and [r["group"] for r in g["rows"]][0] == "DRV-104"  # the largest amount at risk comes first
    m = ex["result"]["metrics"]
    assert sum(r["exceptions"] for r in g["rows"]) == m["exceptions"]
    assert round(sum(r["at_risk"] for r in g["rows"]), 2) == round(m["amount_difference"] + m["unmatched_amount"], 2)
    assert all(r["rows"] == 12 for r in g["rows"])  # 60 orders over 5 drivers


def test_excel_report_lists_matches_and_drivers(client):
    ex = run_pos(client)
    wb = load_workbook(settings.reports_dir / f"execution_{ex['id']}_report.xlsx")
    assert wb.sheetnames == ["Summary", "Exceptions", "By driver", "Matched", "Columns", "Data"]
    assert wb["Matched"].max_row == 1 + 53 and wb["Exceptions"].max_row == 1 + 10
    facts = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert facts["Rows matched"] == 53 and facts["Amount at risk (both added)"] == 1426.31
    assert [c.value for c in wb["By driver"][1]][0] == "Driver"


def test_compare_without_a_group_column_still_works_and_lists_matches():
    left = pd.DataFrame({"id": ["A1", "A2"], "amt": [10, 20]})
    right = pd.DataFrame({"id": ["a1 ", "A2"], "amt": [10, 25]})
    cfg = {"left": "x", "right": "y", "left_key": "id", "right_key": "id", "left_amount": "amt", "right_amount": "amt", "tolerance": 0}
    out, m, d, matched = reconcile.compare_full(left, right, cfg)
    assert "group_by" not in d and m["matched_rows"] == 1 and list(matched["key"]) == ["A1"] and out.iloc[0]["difference"] == -5
    assert reconcile.compare(left, right, cfg)[1] == m  # the older entry point gives the same answer


def test_exceptions_use_the_other_table_when_the_first_has_no_row():
    left = pd.DataFrame({"id": ["A1"], "amt": [10], "driver": ["D1"]})
    right = pd.DataFrame({"id": ["A1", "Z9"], "amt": [10, 50], "driver": ["D1", "D2"]})
    cfg = {"left": "x", "right": "y", "left_key": "id", "right_key": "id", "left_amount": "amt", "right_amount": "amt", "tolerance": 0, "group_by": "driver"}
    rows = reconcile.compare_full(left, right, cfg)[2]["group_by"]["rows"]
    assert rows == [{"group": "D2", "rows": 0, "exceptions": 1, "net_difference": 0.0, "no_partner": 50.0, "at_risk": 50.0}]


def test_the_group_column_is_checked_when_the_workflow_is_saved_and_run(client):
    steps = [{"id": "s1", "type": "compare", "name": "m", "config": {"left": "a", "right": "b", "left_key": "id", "right_key": "id", "group_by": "x" * 80}}]
    assert any("group by" in e for e in clean_definition(steps)[1])
    wf = client.post("/api/templates/pos-reconciliation/use").json()
    full = client.get(f"/api/workflows/{wf['id']}").json()
    steps = full["definition"]["steps"]
    next(s for s in steps if s["type"] == "compare")["config"]["group_by"] = "no_such_column"
    saved = client.put(f"/api/workflows/{wf['id']}", json={"name": full["name"], "description": full["description"], "category": full["category"], "steps": steps})
    assert saved.status_code == 200  # the column is only checked against the real files when the run starts
    files = {"system": {"source": "samples", "file": "pos_system_orders.csv"}, "pos": {"source": "samples", "file": "pos_collections.csv"}}
    ex = client.post(f"/api/workflows/{wf['id']}/run", json={"inputs": files}).json()
    assert ex["status"] == "failed" and "no_such_column" in ex["error_message"]
